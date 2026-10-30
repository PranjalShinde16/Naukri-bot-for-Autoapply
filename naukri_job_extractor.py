
"""Extract Naukri job cards to Excel without opening job URLs."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


# Companies whose job listings should be excluded from the Excel output.
EXCLUDED_COMPANIES = {
    "infosys",
    "tata consultancy services",
    "capgemini",
    "wipro",
    "cgi",
    "cognizant",
    "epam systems",
    "accenture",
    "infosys bpm",
}


def filter_excluded_companies(jobs):
    """Remove jobs whose company name exactly matches the exclusion list."""

    filtered_jobs = []
    excluded_jobs = []

    for job in jobs:
        company = (job.get("company") or "").strip()
        normalized_company = " ".join(company.casefold().split())

        if normalized_company in EXCLUDED_COMPANIES:
            excluded_jobs.append(job)
        else:
            filtered_jobs.append(job)

    print(
        f"\nExcluded-company filter: removed {len(excluded_jobs)} job(s).",
        flush=True,
    )

    for job in excluded_jobs:
        print(
            f"  Excluded: {job.get('company')} — {job.get('title')}",
            flush=True,
        )

    print(
        f"Jobs remaining after exclusion: {len(filtered_jobs)}",
        flush=True,
    )

    return filtered_jobs


def scrape_first_page_jobs_to_excel(page, append: bool = False) -> list[str]:
    """
    Extract job cards from the current Naukri results page.

    When append=False, start a fresh workbook for a new automation run.
    When append=True, append only URLs not already present in the workbook.

    Returns URLs newly added from this page, after company filtering.
    """
    from openpyxl import load_workbook

    output_path = Path(__file__).resolve().parent / "naukri_job_list.xlsx"
    page_url = getattr(page, "url", "<unknown>")

    print(
        f"\n[SCRAPE] Starting page extraction. "
        f"append={append}; page_url={page_url}",
        flush=True,
    )

    try:
        page.wait_for_function(
            """() => document.querySelectorAll(
                'a[href*="/job-listings-"]'
            ).length > 0""",
            timeout=25000,
        )
    except Exception as exc:
        print(
            f"[SCRAPE][ERROR] Job-listing anchors did not appear: "
            f"{type(exc).__name__}: {exc}",
            flush=True,
        )

        try:
            stamp = __import__("datetime").datetime.now().strftime(
                "%Y%m%d_%H%M%S"
            )
            html_path = (
                output_path.parent
                / f"naukri_scrape_failure_{stamp}.html"
            )
            png_path = (
                output_path.parent
                / f"naukri_scrape_failure_{stamp}.png"
            )

            html_path.write_text(page.content(), encoding="utf-8")
            page.screenshot(path=str(png_path), full_page=True)

            print(
                f"[SCRAPE][DIAGNOSTIC] HTML saved: {html_path.name}",
                flush=True,
            )
            print(
                f"[SCRAPE][DIAGNOSTIC] Screenshot saved: {png_path.name}",
                flush=True,
            )

        except Exception as diagnostic_exc:
            print(
                "[SCRAPE][DIAGNOSTIC] Could not save diagnostic artifacts: "
                f"{diagnostic_exc}",
                flush=True,
            )

        raise

    jobs = page.evaluate("""() => {
        const clean = value =>
            (value || '').replace(/\\s+/g, ' ').trim();

        const jobAnchors = [
            ...document.querySelectorAll('a[href*="/job-listings-"]')
        ];

        const seen = new Set();
        const rows = [];

        const companySelectors = [
            'a.comp-name',
            '.comp-name',
            '[class*="comp-name"]',
            '[class*="companyName"]',
            '[class*="company-name"]',
            '[data-testid*="company"]',
            'a[href*="/company/"]',
            'a[href*="/companies/"]'
        ];

        const titleSelectors = [
            'a.title',
            '[class*="jobTitle"]',
            '[class*="title"]',
            '[data-testid*="title"]'
        ];

        const getText = element => clean(
            element && (
                element.innerText ||
                element.textContent ||
                element.getAttribute('title')
            )
        );

        const getCompany = card => {
            if (!card) return '';

            for (const selector of companySelectors) {
                for (const element of card.querySelectorAll(selector)) {
                    const text = getText(element);

                    if (
                        text &&
                        text.length < 160 &&
                        !/^(view details|reviews?|save)$/i.test(text)
                    ) {
                        return text;
                    }
                }
            }

            return '';
        };

        const getTitle = (anchor, card) => {
            let title = getText(anchor);

            if (!title && card) {
                for (const selector of titleSelectors) {
                    const element = card.querySelector(selector);
                    title = getText(element);

                    if (title) break;
                }
            }

            return title;
        };

        for (const anchor of jobAnchors) {
            let url = '';

            try {
                url = new URL(
                    anchor.getAttribute('href'),
                    location.origin
                ).href;
            } catch (_) {
                continue;
            }

            if (!url.includes('/job-listings-')) continue;

            let card = anchor.closest(
                '.srp-jobtuple-wrapper, ' +
                '.cust-job-tuple, [data-job-id], article'
            );

            if (!card) {
                let node = anchor.parentElement;

                for (
                    let depth = 0;
                    node && depth < 9;
                    depth++, node = node.parentElement
                ) {
                    const textLength = clean(
                        node.innerText || ''
                    ).length;

                    if (
                        textLength > 0 &&
                        textLength < 2200 &&
                        node.querySelector(
                            'a[href*="/job-listings-"]'
                        )
                    ) {
                        card = node;
                        break;
                    }
                }
            }

            const title = getTitle(anchor, card);

            if (!title || seen.has(url)) continue;

            const company = getCompany(card);

            seen.add(url);
            rows.push({ company, title, url });
        }

        return {
            anchorCount: jobAnchors.length,
            jobs: rows
        };
    }""")

    raw_jobs = jobs.get("jobs", []) if isinstance(jobs, dict) else jobs
    anchor_count = (
        jobs.get("anchorCount", "n/a")
        if isinstance(jobs, dict)
        else "n/a"
    )

    print(
        f"[SCRAPE] DOM anchors={anchor_count}; "
        f"usable extracted jobs={len(raw_jobs)}",
        flush=True,
    )

    unique_jobs = []
    seen_urls = set()

    for job in raw_jobs:
        url = (job.get("url") or "").strip()

        if not url or url in seen_urls:
            continue

        seen_urls.add(url)

        unique_jobs.append({
            "company": (job.get("company") or "").strip(),
            "title": (job.get("title") or "").strip(),
            "url": url,
        })

    extracted_count = len(unique_jobs)
    filtered_jobs = filter_excluded_companies(unique_jobs)
    excluded_count = extracted_count - len(filtered_jobs)

    print(
        f"[SCRAPE] Extracted={extracted_count}; "
        f"excluded_companies={excluded_count}; "
        f"eligible_this_page={len(filtered_jobs)}",
        flush=True,
    )

    workbook = None

    try:
        if append and output_path.exists():
            workbook = load_workbook(output_path)
            sheet = workbook.active

            headers = [
                str(cell.value or "").strip().lower()
                for cell in sheet[1]
            ]

            expected = ["company name", "job title", "job url"]

            if headers[:3] != expected:
                raise ValueError(
                    f"Unexpected workbook headers in "
                    f"{output_path.name}: {headers!r}"
                )

            print(
                "[SCRAPE][EXCEL] Appending to existing workbook; "
                f"rows_before={sheet.max_row}",
                flush=True,
            )

        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Naukri Jobs"

            sheet.append([
                "Company Name",
                "Job Title",
                "Job URL",
            ])

            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(
                    "solid",
                    fgColor="2457C5",
                )
                cell.alignment = Alignment(vertical="center")

            sheet.freeze_panes = "A2"
            sheet.column_dimensions["A"].width = 32
            sheet.column_dimensions["B"].width = 48
            sheet.column_dimensions["C"].width = 100

            print(
                "[SCRAPE][EXCEL] Started a fresh workbook for this run.",
                flush=True,
            )

        existing_urls = set()

        for row_number in range(2, sheet.max_row + 1):
            value = str(
                sheet.cell(row_number, 3).value or ""
            ).strip()

            if value:
                existing_urls.add(value)

        newly_added_urls = []
        duplicate_count = 0
        missing_company_count = 0

        for job in filtered_jobs:
            url = job["url"]

            if url in existing_urls:
                duplicate_count += 1

                print(
                    "[SCRAPE][DEDUP] Skipping URL already in workbook: "
                    f"{url}",
                    flush=True,
                )
                continue

            if not job["company"]:
                missing_company_count += 1

            sheet.append([
                job["company"],
                job["title"],
                url,
            ])

            url_cell = sheet.cell(
                row=sheet.max_row,
                column=3,
            )
            url_cell.hyperlink = url
            url_cell.style = "Hyperlink"

            existing_urls.add(url)
            newly_added_urls.append(url)

        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 24

        workbook.save(output_path)

        print(
            f"[SCRAPE][EXCEL] Saved {output_path.name}; "
            f"rows_after={sheet.max_row}; "
            f"new_urls={len(newly_added_urls)}; "
            f"duplicates_skipped={duplicate_count}",
            flush=True,
        )

        if missing_company_count:
            print(
                f"[SCRAPE][WARNING] {missing_company_count} job(s) "
                "have no company name.",
                flush=True,
            )

        print(
            "[SCRAPE] Current-page URLs queued for processing: "
            f"{len(newly_added_urls)}",
            flush=True,
        )

        return newly_added_urls

    finally:
        if workbook is not None:
            workbook.close()
